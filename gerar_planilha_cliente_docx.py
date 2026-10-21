"""Gera PLANILHA_CLIENTE.docx — mesmo conteudo do PDF simplificado,
em formato Word pra editar.
"""

from pathlib import Path

from docx import Document
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from docx.shared import Cm, Pt, RGBColor


COR_PRIMARIA = RGBColor(0x1F, 0x3A, 0x68)  # azul Janco
COR_CINZA_ESCURO = RGBColor(0x55, 0x55, 0x55)
COR_HEADER_BG = "1F3A68"
COR_LINHA_PAR = "F4F6FA"


def _set_cell_bg(cell, hex_color: str) -> None:
    """Pinta o fundo de uma celula via XML (python-docx nao expõe)."""
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_color)
    tc_pr.append(shd)


def _runs_bold(cell, texto: str, cor: RGBColor = None, size: int = 11) -> None:
    """Limpa a celula e escreve um texto em negrito."""
    cell.text = ""
    p = cell.paragraphs[0]
    r = p.add_run(texto)
    r.bold = True
    r.font.name = "Arial"
    r.font.size = Pt(size)
    if cor is not None:
        r.font.color.rgb = cor


def _run_normal(cell, texto: str, size: int = 10) -> None:
    cell.text = ""
    p = cell.paragraphs[0]
    r = p.add_run(texto)
    r.font.name = "Arial"
    r.font.size = Pt(size)


def _header_row(table, textos: list[str], size: int = 11) -> None:
    linha = table.rows[0]
    for i, txt in enumerate(textos):
        cell = linha.cells[i]
        _runs_bold(cell, txt, cor=RGBColor(0xFF, 0xFF, 0xFF), size=size)
        _set_cell_bg(cell, COR_HEADER_BG)
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER


def _data_row(table, idx_linha: int, valores: list[str], zebra: bool = False) -> None:
    linha = table.rows[idx_linha]
    for i, txt in enumerate(valores):
        cell = linha.cells[i]
        _run_normal(cell, txt)
        if zebra:
            _set_cell_bg(cell, COR_LINHA_PAR)
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER


def h1(doc: Document, texto: str) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(16)
    p.paragraph_format.space_after = Pt(6)
    r = p.add_run(texto)
    r.font.name = "Arial"
    r.font.size = Pt(15)
    r.bold = True
    r.font.color.rgb = COR_PRIMARIA


def paragrafo(doc: Document, html_texto: str) -> None:
    """Escreve paragrafo tratando <b>...</b> como negrito inline."""
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(8)
    import re
    partes = re.split(r"(<b>.*?</b>)", html_texto)
    for parte in partes:
        if not parte:
            continue
        if parte.startswith("<b>") and parte.endswith("</b>"):
            r = p.add_run(parte[3:-4])
            r.bold = True
        else:
            r = p.add_run(parte)
        r.font.name = "Arial"
        r.font.size = Pt(11)


def destaque(doc: Document, html_texto: str) -> None:
    """Caixa destacada (fundo azul claro)."""
    table = doc.add_table(rows=1, cols=1)
    table.autofit = False
    table.columns[0].width = Cm(17)
    cell = table.rows[0].cells[0]
    cell.text = ""
    _set_cell_bg(cell, "EEF2FA")
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(4)
    import re
    partes = re.split(r"(<b>.*?</b>)", html_texto)
    for parte in partes:
        if not parte:
            continue
        if parte.startswith("<b>") and parte.endswith("</b>"):
            r = p.add_run(parte[3:-4])
            r.bold = True
        else:
            r = p.add_run(parte)
        r.italic = True
        r.font.name = "Arial"
        r.font.size = Pt(10.5)
        r.font.color.rgb = COR_PRIMARIA
    # espaco apos a caixa
    doc.add_paragraph()


def bullet(doc: Document, html_texto: str) -> None:
    p = doc.add_paragraph(style="List Bullet")
    import re
    partes = re.split(r"(<b>.*?</b>)", html_texto)
    for parte in partes:
        if not parte:
            continue
        if parte.startswith("<b>") and parte.endswith("</b>"):
            r = p.add_run(parte[3:-4])
            r.bold = True
        else:
            r = p.add_run(parte)
        r.font.name = "Arial"
        r.font.size = Pt(11)


def tabela_colunas(doc: Document) -> None:
    linhas_dados = [
        ("Data de vencimento", "Sim",
         "Dia em que o pagamento vence"),
        ("Valor", "Sim",
         "Valor em reais do pagamento"),
        ("Data de pagamento", "Recomendado",
         "Dia em que o dinheiro realmente saiu do banco"),
        ("Nº da nota fiscal", "Recomendado",
         "Número da NF associada"),
        ("CNPJ do fornecedor", "Recomendado",
         "CNPJ de quem recebeu o pagamento"),
        ("Nome do fornecedor", "Recomendado",
         "Razão social ou nome fantasia"),
        ("Histórico", "Opcional",
         "Descrição / observação do pagamento"),
        ("Tipo / Categoria", "Opcional",
         "Classificação da despesa (ex: aluguel, luz, material)"),
    ]

    table = doc.add_table(rows=1 + len(linhas_dados), cols=3)
    table.style = "Light Grid Accent 1"
    table.autofit = False
    # Larguras
    larguras = [Cm(5), Cm(3.3), Cm(8.7)]
    for col, larg in zip(table.columns, larguras):
        col.width = larg
    for row in table.rows:
        for cell, larg in zip(row.cells, larguras):
            cell.width = larg

    _header_row(table, ["Coluna", "Precisa ter?", "O que preencher"])

    for i, (col, prec, oq) in enumerate(linhas_dados, start=1):
        linha = table.rows[i]
        _runs_bold(linha.cells[0], col, size=10)
        _run_normal(linha.cells[1], prec)
        linha.cells[1].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        _run_normal(linha.cells[2], oq)
        linha.cells[0].vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        linha.cells[1].vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        linha.cells[2].vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        if i % 2 == 1:
            for cell in linha.cells:
                _set_cell_bg(cell, COR_LINHA_PAR)


def tabela_exemplo(doc: Document) -> None:
    cabecalho = [
        "Vencimento", "Pagamento", "Valor", "NF",
        "CNPJ", "Fornecedor", "Histórico", "Tipo",
    ]
    linhas_dados = [
        ("05/10/2026", "05/10/2026", "1.234,56", "12345",
         "12.345.678/0001-90", "ACME LTDA",
         "Comissão outubro", "Serviço"),
        ("10/10/2026", "11/10/2026", "500,00", "67890",
         "98.765.432/0001-10", "FORNECEDOR X",
         "Material escritório", "Material"),
        ("15/10/2026", "15/10/2026", "2.789,30", "11111",
         "11.222.333/0001-44", "ENERGIA SA",
         "Conta de luz", "Utilidade"),
    ]
    table = doc.add_table(rows=1 + len(linhas_dados), cols=len(cabecalho))
    table.style = "Light Grid Accent 1"
    table.autofit = True

    _header_row(table, cabecalho, size=9)

    for i, dados in enumerate(linhas_dados, start=1):
        linha = table.rows[i]
        for j, txt in enumerate(dados):
            cell = linha.cells[j]
            _run_normal(cell, txt, size=9)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            if i % 2 == 1:
                _set_cell_bg(cell, COR_LINHA_PAR)


def construir(saida: Path) -> None:
    doc = Document()

    # Margens
    for section in doc.sections:
        section.left_margin = Cm(2)
        section.right_margin = Cm(2)
        section.top_margin = Cm(1.6)
        section.bottom_margin = Cm(2)

    # Fonte padrao Arial no estilo Normal
    estilo_normal = doc.styles["Normal"]
    estilo_normal.font.name = "Arial"
    estilo_normal.font.size = Pt(11)

    # Titulo
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(4)
    r = p.add_run("Planilha de controle de pagamentos")
    r.font.name = "Arial"
    r.font.size = Pt(22)
    r.bold = True
    r.font.color.rgb = COR_PRIMARIA

    # Subtitulo
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(20)
    r = p.add_run("Como montar a sua planilha para enviar à contabilidade")
    r.font.name = "Arial"
    r.font.size = Pt(11)
    r.italic = True
    r.font.color.rgb = COR_CINZA_ESCURO

    paragrafo(doc,
        "Esta é a planilha de controle dos pagamentos da sua empresa. "
        "Nós usamos ela para conferir com o extrato bancário e com os "
        "lançamentos do sistema contábil, garantindo que tudo está "
        "batendo certinho.",
    )
    paragrafo(doc,
        "Para funcionar bem, a planilha precisa ter algumas colunas. "
        "Veja abaixo quais são:",
    )

    h1(doc, "Colunas da planilha")
    tabela_colunas(doc)
    doc.add_paragraph()

    destaque(doc,
        "<b>Quanto mais colunas preenchidas, melhor.</b> Com mais "
        "informação, o sistema consegue bater automaticamente mais "
        "pagamentos e você recebe os relatórios mais rápido, com "
        "menos ajustes manuais.",
    )

    h1(doc, "Exemplo de como deve ficar")
    paragrafo(doc,
        "Veja um exemplo de planilha preenchida (os valores são "
        "fictícios, só para ilustrar o formato):",
    )
    tabela_exemplo(doc)
    doc.add_paragraph()

    h1(doc, "Como preencher cada coluna")
    bullet(doc,
        "<b>Datas</b> no formato <b>dd/mm/aaaa</b> (ex: 05/10/2026) ou "
        "formatadas como data no Excel.",
    )
    bullet(doc,
        "<b>Valor</b> em reais, com vírgula para centavos "
        "(ex: 1.234,56). Não precisa colocar o símbolo R$.",
    )
    bullet(doc,
        "<b>Nº da nota fiscal</b>: só o número, sem letras ou pontos.",
    )
    bullet(doc,
        "<b>CNPJ</b>: pode preencher com ou sem máscara "
        "(12.345.678/0001-90 ou 12345678000190 — os dois funcionam).",
    )
    bullet(doc,
        "<b>Fornecedor</b>: nome da empresa que recebeu o pagamento. "
        "A razão social é o ideal, mas o nome fantasia também vale.",
    )

    h1(doc, "Dicas importantes")
    bullet(doc,
        "<b>A primeira linha da planilha é o cabeçalho</b> — ponha os "
        "nomes das colunas na linha 1 e os pagamentos a partir da "
        "linha 2.",
    )
    bullet(doc,
        "<b>Uma planilha por empresa.</b> Se a sua empresa tem matriz "
        "e filiais, faça uma planilha separada para cada uma.",
    )
    bullet(doc,
        "<b>Evite células mescladas</b> (aquelas que juntam várias "
        "colunas numa só) nas linhas de pagamento. Pode usar mescla "
        "só no título acima do cabeçalho, se quiser.",
    )
    bullet(doc,
        "<b>Valor em uma coluna só.</b> Se você hoje separa pagamento "
        "em colunas de \"débito\" e \"crédito\", prefira juntar numa "
        "coluna só \"Valor\".",
    )
    bullet(doc,
        "<b>Linhas em branco</b> no meio da planilha não atrapalham, "
        "podem ficar.",
    )
    bullet(doc,
        "<b>Salve como .xlsx</b> (formato padrão do Excel).",
    )

    h1(doc, "E se faltar alguma informação?")
    paragrafo(doc,
        "Tudo bem — a planilha funciona mesmo que você só tenha "
        "<b>vencimento e valor</b>. As outras colunas são para melhorar "
        "a conferência automática. Se você já tem uma planilha pronta "
        "com nomes de coluna um pouco diferentes dos nomes aqui "
        "(ex: \"Val Pgto\" em vez de \"Valor\"), também não tem problema: "
        "o sistema deixa a gente ajustar isso na importação. O "
        "importante é ter as informações.",
    )

    doc.add_paragraph()
    destaque(doc,
        "Qualquer dúvida sobre como montar a planilha, pode falar com "
        "a gente — a Janco está à disposição.",
    )

    doc.save(str(saida))


def main() -> None:
    saida = Path(__file__).parent / "PLANILHA_CLIENTE.docx"
    construir(saida)
    print(f"DOCX gerado: {saida}")


if __name__ == "__main__":
    main()
